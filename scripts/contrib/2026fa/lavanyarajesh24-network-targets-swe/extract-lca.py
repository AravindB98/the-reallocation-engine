#!/usr/bin/env python3
"""extract-lca.py — SOC-coded H-1B evidence from DOL OFLC LCA disclosure data (closes recipe TODO 1).

OFFLINE. Reads one DOL LCA disclosure workbook that a person downloaded by hand
(dol.gov blocks scripted downloads) and writes a small per-employer aggregate:

  certified H-1B LCAs whose SOC code is 15-1252 (Software Developers), counted
  nationwide and, separately, at worksites in the persona's target states, with
  received-date range and wage-level counts.

Nationwide is the policy signal ("does this employer sponsor software developers at
all?"); the target-state count is context. The first version counted target states
only and wrongly treated employers whose software hires sit at remote worksites in
other states as non-sponsors — see the worked run.

    python3 scripts/contrib/2026fa/lavanyarajesh24-network-targets-swe/extract-lca.py \
        ~/Downloads/LCA_Disclosure_Data_FY2026_Q3.xlsx \
        [--persona inputs/persona.json] [--out data/lca-15-1252.csv]

Writes <out> and <out>.provenance.json (source file name, size, sha256, rows read,
received-date range actually seen). The raw workbook is NOT copied into the repo.

What an LCA is not: a certified Labor Condition Application is the employer's wage
attestation filed before an H-1B petition. It is not a petition approval, it covers
extensions and transfers as well as new hires, and one LCA can cover several workers.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import statistics
import sys
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent

REQUIRED = ["CASE_STATUS", "RECEIVED_DATE", "VISA_CLASS", "SOC_CODE", "EMPLOYER_NAME", "WORKSITE_STATE",
            "WAGE_RATE_OF_PAY_FROM", "WAGE_UNIT_OF_PAY", "PW_WAGE_LEVEL"]
ANNUAL = {"Year": 1, "Month": 12, "Bi-Weekly": 26, "Week": 52, "Hour": 2080}

# Must match nameKey() in lib.mjs exactly — the join is exact-key only.
SUFFIXES = ["incorporated", "corporation", "company", "limited", "inc", "llc", "corp", "ltd", "co", "lp", "pbc"]


def name_key(name: str) -> str:
    k = re.sub(r"[^a-z0-9]", "", str(name or "").lower())
    changed = True
    while changed:
        changed = False
        for s in SUFFIXES:
            if len(k) > len(s) and k.endswith(s):
                k = k[: -len(s)]
                changed = True
                break
    return k


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def as_day(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            pass
    return None


def annual_wage(rate, unit) -> float | None:
    try:
        r = float(str(rate).replace(",", "").replace("$", ""))
    except (TypeError, ValueError):
        return None
    m = ANNUAL.get(str(unit or "").strip())
    return r * m if m else None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--persona", type=Path, default=HERE / "inputs/persona.json")
    ap.add_argument("--out", type=Path, default=HERE / "data/lca-15-1252.csv")
    ap.add_argument("--soc", default=None, help="default: persona.target.soc_code")
    a = ap.parse_args()

    import openpyxl  # imported late so --help works without it

    persona = json.loads(a.persona.read_text())
    soc = a.soc or persona["target"]["soc_code"]
    states = {s.upper() for s in persona["target"]["states"]}

    if not a.workbook.exists():
        print(f"✗ workbook not found: {a.workbook}", file=sys.stderr)
        return 2
    wb = openpyxl.load_workbook(a.workbook, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        print(f"✗ workbook lacks required column(s): {', '.join(missing)}", file=sys.stderr)
        return 2
    ix = {c: header.index(c) for c in header if c}

    read = kept = 0
    seen_min = seen_max = kept_min = kept_max = None
    agg: dict[str, dict] = {}
    for r in rows:
        read += 1
        if read % 100_000 == 0:
            print(f"  … {read:,} rows read, {kept:,} kept", flush=True)
        day = as_day(r[ix["RECEIVED_DATE"]])
        if day:
            seen_min = day if seen_min is None or day < seen_min else seen_min
            seen_max = day if seen_max is None or day > seen_max else seen_max
        if str(r[ix["VISA_CLASS"]] or "").strip() != "H-1B":
            continue
        if str(r[ix["CASE_STATUS"]] or "").strip() != "Certified":
            continue
        if not str(r[ix["SOC_CODE"]] or "").strip().startswith(soc):
            continue
        st = str(r[ix["WORKSITE_STATE"]] or "").strip().upper()
        in_target = st in states
        kept += 1
        if day:
            kept_min = day if kept_min is None or day < kept_min else kept_min
            kept_max = day if kept_max is None or day > kept_max else kept_max
        name = str(r[ix["EMPLOYER_NAME"]] or "").strip()
        k = name_key(name)
        e = agg.setdefault(k, {"names": set(), "n": 0, "n_target": 0, "levels": {}, "wages": [], "first": None, "last": None, "states": set()})
        e["names"].add(name)
        e["n"] += 1
        e["n_target"] += in_target
        lvl = str(r[ix["PW_WAGE_LEVEL"]] or "").strip() or "none"
        e["levels"][lvl] = e["levels"].get(lvl, 0) + 1
        w = annual_wage(r[ix["WAGE_RATE_OF_PAY_FROM"]], r[ix["WAGE_UNIT_OF_PAY"]])
        if w:
            e["wages"].append(w)
        e["states"].add(st)
        if day:
            e["first"] = day if e["first"] is None or day < e["first"] else e["first"]
            e["last"] = day if e["last"] is None or day > e["last"] else e["last"]

    a.out.parent.mkdir(parents=True, exist_ok=True)
    cols = ["employer_key", "employer_names_seen", "certified_lcas", "certified_in_target_states", "worksite_states", "level_I", "level_II", "level_III", "level_IV",
            "level_none", "median_annual_wage_offered", "first_received", "last_received"]
    with a.out.open("w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f)
        wr.writerow(cols)
        for k in sorted(agg, key=lambda x: (-agg[x]["n"], x)):
            e = agg[k]
            lv = e["levels"]
            wr.writerow([k, " / ".join(sorted(e["names"]))[:300], e["n"], e["n_target"], ";".join(sorted(e["states"])),
                         lv.get("I", 0), lv.get("II", 0), lv.get("III", 0), lv.get("IV", 0), lv.get("none", 0),
                         round(statistics.median(e["wages"])) if e["wages"] else "", e["first"] or "", e["last"] or ""])

    prov = {
        "_tool": "extract-lca.py",
        "source": "U.S. Department of Labor, OFLC, LCA Programs (H-1B, H-1B1, E-3) disclosure data",
        "source_page": "https://www.dol.gov/agencies/eta/foreign-labor/performance",
        "source_file": a.workbook.name,
        "source_bytes": a.workbook.stat().st_size,
        "source_sha256": sha256(a.workbook),
        "filter": {"VISA_CLASS": "H-1B", "CASE_STATUS": "Certified", "SOC_CODE_prefix": soc, "target_states_counted_separately": sorted(states)},
        "rows_read": read,
        "rows_kept": kept,
        "rows_kept_in_target_states": sum(e["n_target"] for e in agg.values()),
        "employers": len(agg),
        "received_date_range_seen": [seen_min, seen_max],
        "received_date_range_kept": [kept_min, kept_max],
        "extracted_at": datetime.now(timezone.utc).isoformat(),
        "not_this": "A certified LCA is a wage attestation, not an H-1B approval; it includes extensions/transfers; one LCA can cover several workers.",
    }
    Path(str(a.out) + ".provenance.json").write_text(json.dumps(prov, indent=2) + "\n")
    print(f"✓ {read:,} rows read → {kept:,} certified H-1B LCAs for SOC {soc} nationwide across {len(agg):,} employers "
          f"({sum(e['n_target'] for e in agg.values()):,} at worksites in {', '.join(sorted(states))})")
    print(f"  received dates: all rows {seen_min} … {seen_max}; kept rows {kept_min} … {kept_max}")
    print(f"  {a.out}  +  {a.out}.provenance.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
